"""
1번 스크립트로 뽑은 서버 로그에서 웹 애플리케이션 로그(p6spy 줄을 뺀 로그)를 엑셀 파일 하나로 정리한다.

시트 구성
  1_원본       : 웹 애플리케이션 로그의 기록 하나를 한 행으로 두고, 시각·로그 레벨·서버 스레드·클래스·메시지로 나눈다.
  2_요청별정리 : 기록을 요청(요청 A, 요청 B, A-1, B-1)별로 묶고, 메시지에서 주문 ID 같은 값을 꺼내 열로 둔다.
  3_시간값     : 요청 A와 요청 B의 주요 시점을 클라이언트 시각과 서버 시각으로 나누어 둔다. 오른쪽에 계산값 표를 둔다.
각 시트 아래에는 열마다 무슨 데이터인지 적은 열 설명 표를 둔다.

입력 파일은 읽기만 한다. 출력 파일이 이미 있으면 덮어쓰지 않고 중단한다.
실행 방법은 같은 폴더의 2_웹애플리케이션로그_엑셀정리_실행방법.txt를 참고한다.
"""
import argparse
import csv
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KST = timezone(timedelta(hours=9))

# 서버 로그 줄의 형식
# 2026-10-15T15:20:44.866+09:00  INFO 650842 --- [ChatService] [http-nio-8186-exec-5] c.c.marketplace.testsupport.TestPause    : 메시지
LOG_LINE = re.compile(
    r'^(?P<time>\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\.\d{3}[+-]\d{2}:\d{2})\s+'
    r'(?P<level>[A-Z]+)\s+\d+\s+---\s+\[[^\]]*\]\s+\[\s*(?P<thread>[^\]]*?)\s*\]\s+'
    r'(?P<logger>\S+)\s*:\s(?P<message>.*)$'
)

# 메시지 문장에서 값을 꺼낼 때 쓰는 패턴
REQUEST_START = 'JWTAuthenticationFilter - 요청 처리 시작'
HTTP_METHOD = re.compile(r'HTTP Method: (?P<method>\S+), URL: \S+?/ChatService(?P<path>\S*)')
AUTH_USER_ID = re.compile(r'JWT 인증 성공: 사용자 ID = (?P<member>\S+)')
ORDER_ID = re.compile(r'orderId=(\d+)|/orders/(\d+)/')
CANCELLED_BY = re.compile(r'cancelledBy=(\w+)')
AMOUNT = re.compile(r'amount=(-?\d+)')
BALANCE_AFTER = re.compile(r'balanceAfter=(-?\d+)')
PAUSE_POINT = re.compile(r'point=(\w+)')

# 메시지의 앞부분으로 기록 종류를 정한다. 위에서부터 차례로 확인한다.
RECORD_KINDS = [
    (REQUEST_START, '요청 처리 시작'),
    ('HTTP Method:', '요청 메서드·URL 기록'),
    ('JWT 인증 성공: 사용자 ID', '인증 성공(사용자 ID)'),
    ('JWT 인증 성공: 사용자 이름', '인증 성공(사용자 이름)'),
    ('[테스트 일시 중지] 대기 시작', '일시 중지 대기 시작'),
    ('[테스트 일시 중지] 재개', '일시 중지 재개'),
    ('[잔액 증가]', '잔액 증가'),
    ('[주문 취소]', '주문 취소'),
]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
TITLE_FONT = Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical='top')


# ---------------------------------------------------------------------
# 입력 파일 읽기
# ---------------------------------------------------------------------
def read_web_records(log_path):
    """1번 추출 파일을 읽어, p6spy 기록을 뺀 웹 애플리케이션 로그 기록 목록을 돌려준다.

    기록 하나는 시각으로 시작하는 줄과 그 뒤의 시각 없는 줄로 이루어진다. p6spy 기록은 딸린 SQL 줄까지 함께 뺀다.
    """
    lines = log_path.read_text(encoding='utf-8').splitlines()
    records = []
    current = None
    for number, line in enumerate(lines, start=1):
        match = LOG_LINE.match(line)
        if match:
            current = {
                'line': number,
                'time': datetime.fromisoformat(match.group('time')),
                'level': match.group('level'),
                'thread': match.group('thread'),
                'logger': match.group('logger'),
                'message': match.group('message'),
            }
            records.append(current)
        elif current is not None and line.strip():
            # 시각 없는 줄은 앞 기록에 딸린 줄이다. 웹 애플리케이션 기록이면 메시지 아래에 붙인다.
            current['message'] += '\n' + line
    return [record for record in records if record['logger'] != 'p6spy']


def read_jtl(jtl_path):
    """JMeter 결과 파일(결과.jtl)을 읽어 요청 이름(label)별 보낸 시각과 걸린 시간을 돌려준다."""
    samples = {}
    with jtl_path.open(encoding='utf-8', newline='') as file:
        for row in csv.DictReader(file):
            sent = datetime.fromtimestamp(int(row['timeStamp']) / 1000, KST)
            samples[row['label']] = {'sent': sent, 'elapsed_ms': int(row['elapsed']), 'code': row['responseCode']}
    return samples


def read_resume_log(resume_path, client_date):
    """재개 제어 기록(재개제어_기록.txt)을 읽어 (시각 원문, 시각, 내용) 목록을 돌려준다. 이 파일의 시각에는 날짜가 없어 client_date를 붙인다."""
    entries = []
    for line in resume_path.read_text(encoding='utf-8').splitlines():
        if not line.strip():
            continue
        time_text, content = line.split(' ', 1)
        hms, _, fraction = time_text.partition('.')
        # 이 파일의 시각은 소수점 아래가 최대 9자리(나노초)이다. 파이썬 datetime은 마이크로초(6자리)까지만 다루므로
        # 계산에 쓰는 값(time)은 6자리로 자르고, 엑셀에 적는 시각 원문(text)은 자르지 않는다.
        moment = datetime.fromisoformat(f'{client_date}T{hms}.{(fraction + "000000")[:6]}').replace(tzinfo=KST)
        entries.append({'text': f'{client_date} {time_text}', 'time': moment, 'content': content})
    return entries


# ---------------------------------------------------------------------
# 기록 정리
# ---------------------------------------------------------------------
def record_kind(message):
    """메시지의 앞부분으로 기록 종류를 정한다. RECORD_KINDS에 없는 메시지는 「그 밖의 기록」이다."""
    for prefix, kind in RECORD_KINDS:
        if message.startswith(prefix):
            return kind
    return '그 밖의 기록'


def first_group(pattern, text):
    """패턴에 맞는 첫 값을 돌려준다. 맞는 값이 없으면 빈 문자열을 돌려주어 엑셀 칸을 비워 둔다."""
    match = pattern.search(text)
    if not match:
        return ''
    return next(group for group in match.groups() if group is not None)


def split_requests(records):
    """기록을 요청 단위로 나눈다.

    서버는 스레드를 여러 요청에 다시 쓰므로, 같은 스레드의 「요청 처리 시작」 기록부터 그 스레드의 다음 「요청 처리 시작」 기록 전까지를 한 요청으로 본다.
    """
    requests = []
    open_request = {}   # 스레드 이름 → 그 스레드에서 지금 처리 중인 요청
    for record in records:
        thread = record['thread']
        if record['message'].startswith(REQUEST_START):
            open_request[thread] = {'thread': thread, 'records': [], 'method': '', 'path': '', 'member': ''}
            requests.append(open_request[thread])
        request = open_request.get(thread)
        if request is None:
            continue
        request['records'].append(record)
        method = HTTP_METHOD.search(record['message'])
        if method and not request['method']:
            request['method'], request['path'] = method.group('method'), method.group('path')
        member = AUTH_USER_ID.search(record['message'])
        if member and not request['member']:
            request['member'] = member.group('member')
    return requests


def label_requests(requests):
    """요청마다 테스트에서 쓰는 이름(요청 A, 요청 B, A-1, B-1)을 붙인다.

    구매 목록 조회를 보낸 회원을 구매자, 판매 목록 조회를 보낸 회원을 판매자로 보고,
    주문 취소 요청은 구매자가 보냈으면 요청 A, 판매자가 보냈으면 요청 B로 정한다.
    """
    buyer = next((r['member'] for r in requests if r['path'].endswith('/orders/purchases')), '')
    seller = next((r['member'] for r in requests if r['path'].endswith('/orders/sales')), '')
    for request in requests:
        path = request['path']
        if path.endswith('/orders/purchases'):
            request['label'] = 'A-1(구매 목록 조회)'
        elif path.endswith('/orders/sales'):
            request['label'] = 'B-1(판매 목록 조회)'
        elif path.endswith('/cancel') and request['member'] == buyer:
            request['label'] = '요청 A(구매자 취소)'
        elif path.endswith('/cancel') and request['member'] == seller:
            request['label'] = '요청 B(판매자 취소)'
        else:
            request['label'] = '그 밖의 요청'
    order = ['요청 A(구매자 취소)', '요청 B(판매자 취소)', 'A-1(구매 목록 조회)', 'B-1(판매 목록 조회)', '그 밖의 요청']
    return sorted(requests, key=lambda r: (order.index(r['label']), r['records'][0]['line']))


def server_text(moment):
    """시각을 서버 로그와 같은 자릿수(밀리초까지)의 문자열로 바꾼다."""
    return moment.strftime('%Y-%m-%d %H:%M:%S.') + f'{moment.microsecond // 1000:03d}'


def find_record(request, kind):
    """요청 안에서 주어진 종류의 첫 기록을 돌려준다. 없으면 None을 돌려준다."""
    return next((record for record in request['records'] if record_kind(record['message']) == kind), None)


# ---------------------------------------------------------------------
# 엑셀 작성
# ---------------------------------------------------------------------
def write_table(sheet, start_row, start_col, headers, rows, title=None):
    """표 하나를 쓰고, 표 다음 빈 행 번호를 돌려준다."""
    row = start_row
    if title:
        sheet.cell(row=row, column=start_col, value=title).font = TITLE_FONT
        row += 1
    for offset, header in enumerate(headers):
        cell = sheet.cell(row=row, column=start_col + offset, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    for values in rows:
        row += 1
        for offset, value in enumerate(values):
            sheet.cell(row=row, column=start_col + offset, value=value).alignment = WRAP
    return row + 2


def set_widths(sheet, widths, start_col=1):
    """start_col 열부터 차례로 열 너비를 정한다."""
    for offset, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(start_col + offset)].width = width


def freeze_header(sheet, cell):
    """cell의 위쪽 행과 왼쪽 열을 고정하고, 처음 선택된 셀을 cell로 둔다.

    선택된 셀이 고정한 영역(A1)에 있으면 엑셀이 움직이는 영역을 1행부터 보여 주어, 머리글 행이 두 번 보인다.
    """
    sheet.freeze_panes = cell
    for selection in sheet.sheet_view.selection:
        if selection.pane == sheet.sheet_view.pane.activePane:
            selection.activeCell = cell
            selection.sqref = cell


def build_sheet_original(workbook, records, log_name):
    """1_원본 시트: 웹 애플리케이션 로그 기록 하나를 한 행으로 둔다."""
    sheet = workbook.active
    sheet.title = '1_원본'
    headers = ['원본 줄 번호', '서버 시각', '로그 레벨', '서버 스레드', '로그를 남긴 클래스', '메시지']
    rows = [[r['line'], server_text(r['time']), r['level'], r['thread'], r['logger'], r['message']] for r in records]
    next_row = write_table(sheet, 1, 1, headers, rows)
    freeze_header(sheet, 'A2')
    set_widths(sheet, [12, 25, 10, 22, 45, 90])
    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['원본 줄 번호', f'{log_name}에서 이 기록이 시작되는 줄의 번호'],
        ['서버 시각', '서버가 로그 줄 맨 앞에 기록한 시각(한국 시각, 밀리초까지)'],
        ['로그 레벨', '로그의 중요도. INFO는 일반 기록, WARN은 경고, ERROR는 오류'],
        ['서버 스레드', '이 기록을 남긴 서버 스레드의 이름. 요청 하나는 서버 스레드 하나가 처리한다'],
        ['로그를 남긴 클래스', '이 기록을 남긴 코드의 클래스 이름. 패키지 이름은 서버 로그에서 앞 글자로 줄여 기록된다'],
        ['메시지', '로그 메시지 원문'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1,
               value='이 시트에는 p6spy 기록(SQL 실행 로그)을 넣지 않았다. p6spy 기록은 DB 측 로그로 따로 정리한다.')


def build_sheet_requests(workbook, requests):
    """2_요청별정리 시트: 기록을 요청별로 묶고, 메시지에서 꺼낸 값을 열로 둔다."""
    sheet = workbook.create_sheet('2_요청별정리')
    headers = ['요청 구분', '요청 주체', '요청 메서드·URL', '서버 스레드', '요청별 기록 순서', '원본 줄 번호', '서버 시각',
               '기록 종류', '주문 ID', '취소 주체', '금액', '변동 후 잔액', '중지 위치', '메시지']
    rows = []
    for request in requests:
        for sequence, record in enumerate(request['records'], start=1):
            message = record['message']
            # 요청 구분·요청 주체·요청 메서드·URL·서버 스레드는 요청 하나에 대한 값이므로 그 요청의 첫 행에만 적는다.
            # 모든 행에 적으면 같은 기록이 여러 번 들어간 것처럼 보인다.
            if sequence == 1:
                request_values = [request['label'], request['member'], f"{request['method']} {request['path']}", request['thread']]
            else:
                request_values = ['', '', '', '']
            rows.append(request_values + [
                sequence, record['line'], server_text(record['time']), record_kind(message),
                first_group(ORDER_ID, message), first_group(CANCELLED_BY, message), first_group(AMOUNT, message),
                first_group(BALANCE_AFTER, message), first_group(PAUSE_POINT, message), message,
            ])
    next_row = write_table(sheet, 1, 1, headers, rows)
    freeze_header(sheet, 'B2')
    set_widths(sheet, [20, 16, 30, 22, 8, 8, 25, 22, 8, 10, 10, 12, 12, 80])
    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['요청 구분', '테스트에서 쓰는 요청 이름. 요청 A는 구매자의 주문 취소, 요청 B는 판매자의 주문 취소, A-1과 B-1은 취소 전에 주문 ID를 얻으려고 보낸 구매·판매 목록 조회'],
        ['요청 주체', '「JWT 인증 성공: 사용자 ID」 기록에 나온, 그 요청을 보낸 회원 ID'],
        ['요청 메서드·URL', '「HTTP Method」 기록에 나온 요청 메서드와 경로. 서버 주소와 /ChatService는 빼고 적었다'],
        ['서버 스레드', '그 요청을 처리한 서버 스레드의 이름'],
        ['요청별 기록 순서', '같은 요청이 남긴 기록 가운데 몇 번째 기록인지를 나타낸다. 서버 시각이 같은 밀리초인 기록들의 앞뒤를 구분할 때 쓴다'],
        ['원본 줄 번호', '1_원본 시트의 원본 줄 번호와 같은 값'],
        ['서버 시각', '서버가 기록한 시각(한국 시각, 밀리초까지)'],
        ['기록 종류', '메시지 앞부분으로 나눈 기록의 종류. 해당하는 종류가 없으면 「그 밖의 기록」이다'],
        ['주문 ID', '메시지의 orderId 값 또는 URL의 주문 번호'],
        ['취소 주체', '「주문 취소」 기록의 cancelledBy 값. BUYER는 구매자, SELLER는 판매자'],
        ['금액', '「잔액 증가」 기록의 amount 값(원)'],
        ['변동 후 잔액', '「잔액 증가」 기록의 balanceAfter 값(원)'],
        ['중지 위치', '「일시 중지 대기 시작」과 「일시 중지 재개」 기록의 point 값. ORDER_READ는 주문 조회 직후, WALLET_READ는 구매자 지갑 조회 직후'],
        ['메시지', '로그 메시지 원문'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1, value=(
        '요청의 범위: 같은 서버 스레드의 「요청 처리 시작」 기록부터 그 스레드의 다음 「요청 처리 시작」 기록 전까지를 한 요청으로 보았다. '
        '서버가 스레드를 여러 요청에 다시 쓰기 때문이다. '
        '요청 구분, 요청 주체, 요청 메서드·URL, 서버 스레드는 요청 하나에 대한 값이므로 각 요청의 첫 행에만 적었다. '
        '그 아래 빈칸은 바로 위에 적힌 요청과 같은 값이다.'))


def build_sheet_times(workbook, requests, samples, resume_entries, jtl_name, resume_name, log_name):
    """3_시간값 시트: 요청 A와 요청 B의 주요 시점을 클라이언트 시각과 서버 시각으로 나누어 두고, 오른쪽에 계산값 표를 둔다."""
    sheet = workbook.create_sheet('3_시간값')
    by_label = {r['label']: r for r in requests}
    rows = []
    calculations = []

    # 공통: 재개 제어가 두 요청의 일시 중지를 확인한 시각
    for entry in resume_entries:
        if entry['content'].startswith('요청 A와 요청 B가 모두 멈춤'):
            rows.append(['공통', '재개 제어가 두 요청의 일시 중지를 확인함', entry['text'], '', resume_name])

    for name, label, sample_prefix in [('A', '요청 A(구매자 취소)', 'A-2.'), ('B', '요청 B(판매자 취소)', 'B-2.')]:
        request = by_label.get(label)
        sample = next((s for key, s in samples.items() if key.startswith(sample_prefix)), None)
        resume_sent = next((e for e in resume_entries if e['content'].startswith(f'요청 {name} 재개 요청')), None)
        resume_done = next((e for e in resume_entries if e['content'].startswith(f'요청 {name} 응답 받음')), None)
        start = find_record(request, '요청 처리 시작') if request else None
        pause = find_record(request, '일시 중지 대기 시작') if request else None
        resumed = find_record(request, '일시 중지 재개') if request else None
        cancelled = find_record(request, '주문 취소') if request else None
        received = sample['sent'] + timedelta(milliseconds=sample['elapsed_ms']) if sample else None

        def client(moment):
            return moment.strftime('%Y-%m-%d %H:%M:%S.') + f'{moment.microsecond // 1000:03d}'

        rows += [
            [label, 'JMeter가 취소 요청을 보냄', client(sample['sent']) if sample else '', '', jtl_name],
            [label, '서버가 요청 처리를 시작함', '', server_text(start['time']) if start else '', log_name],
            [label, '서버가 중지 위치에서 대기를 시작함', '', server_text(pause['time']) if pause else '', log_name],
            [label, '재개 제어가 재개 요청의 응답을 받음', resume_sent['text'] if resume_sent else '', '', resume_name],
            [label, '서버가 대기를 끝내고 처리를 이어 감', '', server_text(resumed['time']) if resumed else '', log_name],
            [label, '서버가 주문 취소 기록을 남김(커밋 전)', '', server_text(cancelled['time']) if cancelled else '', log_name],
            [label, 'JMeter가 취소 요청의 응답을 받음(계산값)', client(received) if received else '', '', jtl_name],
            [label, '재개 제어가 취소 요청의 응답 수신을 확인함', resume_done['text'] if resume_done else '', '', resume_name],
        ]

        if pause and resumed:
            wait_ms = (resumed['time'] - pause['time']).total_seconds() * 1000
            calculations.append([f'{label}의 대기 시간',
                                 '「서버가 대기를 끝내고 처리를 이어 감」의 서버 시각 − 「서버가 중지 위치에서 대기를 시작함」의 서버 시각',
                                 f'{wait_ms:.0f} ms', '요청이 서버의 중지 위치에서 멈춰 있던 시간이다. 두 시각 모두 서버 시각이다'])
        if resumed and resume_sent:
            skew_ms = (resumed['time'] - resume_sent['time']).total_seconds() * 1000
            calculations.append([f'{label}에서 잰 두 시계의 차이',
                                 '「서버가 대기를 끝내고 처리를 이어 감」의 서버 시각 − 「재개 제어가 재개 요청의 응답을 받음」의 클라이언트 시각',
                                 f'{skew_ms:.1f} ms',
                                 '서버 시계가 클라이언트 시계보다 앞선 정도이다. 두 사건은 거의 같은 순간에 일어나지만, '
                                 '재개 요청의 응답이 클라이언트에 도착하는 데 걸린 시간이 섞여 있어 실제 차이와 몇 밀리초 다를 수 있다'])
        if received:
            calculations.append([f'{label}의 응답을 받은 시각', f'{jtl_name}의 timeStamp + elapsed',
                                 client(received),
                                 f'{jtl_name}의 timeStamp는 요청을 보낸 시각이고, elapsed는 응답을 받기까지 걸린 시간(ms)이다. '
                                 '응답을 받은 시각은 원본에 따로 기록되지 않아 이 둘을 더해 구했다'])

    headers = ['요청 구분', '시점', '클라이언트 시각', '서버 시각', '출처 파일']
    next_row = write_table(sheet, 1, 1, headers, rows)
    freeze_header(sheet, 'A2')
    set_widths(sheet, [20, 40, 31, 25, 18])
    write_table(sheet, 1, 7, ['계산값', '계산식', '값', '설명'], calculations, title=None)
    sheet.cell(row=len(calculations) + 3, column=7, value='위 표의 값은 원본 파일에 없고, 왼쪽 표의 시각으로 계산한 값이다.')
    set_widths(sheet, [30, 50, 24, 60], start_col=7)
    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['요청 구분', '요청 A(구매자 취소), 요청 B(판매자 취소). 「공통」은 두 요청에 함께 해당하는 시점이다'],
        ['시점', '그 행이 나타내는 사건과 그 사건을 기록한 쪽. 「JMeter」는 취소 요청을 보낸 JMeter 스레드, '
               '「재개 제어」는 일시 중지된 두 요청을 차례로 재개시키는 JMeter 스레드, 「서버」는 미니PC의 웹 애플리케이션이다. '
               'JMeter와 재개 제어가 기록한 시각은 클라이언트 시각 열에, 서버가 기록한 시각은 서버 시각 열에 있다'],
        ['클라이언트 시각', f'JMeter를 실행한 노트북의 시계로 기록한 시각(한국 시각). {jtl_name}, {resume_name}에서 가져왔다. '
                     f'{resume_name}에는 날짜가 없어 {jtl_name}의 날짜를 붙였다'],
        ['서버 시각', f'미니PC의 시계로 기록한 시각(한국 시각). {log_name}에서 가져왔다'],
        ['출처 파일', '그 시각을 가져온 파일'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1, value=(
        '클라이언트 시계와 서버 시계는 서로 맞춰져 있지 않으므로, 클라이언트 시각과 서버 시각을 서로 빼거나 앞뒤를 비교하지 않는다. '
        '하나만 써야 할 때는 서버 시각을 쓴다. 두 시계의 차이는 오른쪽 계산값 표에 있다. '
        '행은 시각순이 아니라 요청별로 사건이 일어나는 순서대로 두었다.'))


def main():
    parser = argparse.ArgumentParser(description='웹 애플리케이션 로그를 엑셀 파일로 정리한다.')
    parser.add_argument('--log', required=True, help='1번 스크립트로 뽑은 서버 로그 파일')
    parser.add_argument('--jtl', required=True, help='JMeter 결과 파일(결과.jtl)')
    parser.add_argument('--resume', required=True, help='재개 제어 기록 파일(재개제어_기록.txt)')
    parser.add_argument('--output', required=True, help='만들 엑셀 파일. 같은 이름의 파일이 있으면 중단한다')
    args = parser.parse_args()

    log_path, jtl_path, resume_path, output = Path(args.log), Path(args.jtl), Path(args.resume), Path(args.output)
    if output.exists():
        sys.exit(f'출력 파일이 이미 있습니다. 다른 이름을 지정하세요: {output}')

    records = read_web_records(log_path)
    requests = label_requests(split_requests(records))
    samples = read_jtl(jtl_path)
    # 재개 제어 기록에는 날짜가 없으므로, 같은 실행의 결과.jtl에 있는 날짜(클라이언트 시각)를 붙인다.
    client_date = min(sample['sent'] for sample in samples.values()).date().isoformat()
    resume_entries = read_resume_log(resume_path, client_date)

    workbook = Workbook()
    build_sheet_original(workbook, records, log_path.name)
    build_sheet_requests(workbook, requests)
    build_sheet_times(workbook, requests, samples, resume_entries, jtl_path.name, resume_path.name, log_path.name)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    print(f'웹 애플리케이션 로그 기록: {len(records)}건')
    for request in requests:
        print(f"  {request['label']}: {request['member']}, {request['method']} {request['path']}, "
              f"{request['thread']}, 기록 {len(request['records'])}건")
    print(f'저장한 파일: {output}')


if __name__ == '__main__':
    main()
