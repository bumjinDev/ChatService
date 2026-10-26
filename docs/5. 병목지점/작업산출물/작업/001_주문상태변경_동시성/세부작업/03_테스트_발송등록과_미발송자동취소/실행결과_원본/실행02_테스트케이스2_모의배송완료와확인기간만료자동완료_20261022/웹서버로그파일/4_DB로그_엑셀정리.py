"""
3번 스크립트로 뽑은 p6spy 기록(서버가 DB에 보낸 SQL과 커밋의 기록)을 엑셀 파일 하나로 정리한다.
이 스크립트는 세부 작업 02의 실행03에서 쓴 같은 이름 스크립트의 사본이다.
이 세부 작업의 실행02(테스트케이스2 주문의 모의 배송 완료와 확인 기간 만료 자동 완료)는 요청 구성이 같아 처리 규칙을 바꾸지 않았다.

시트 구성
  1_원본       : p6spy 기록 하나를 한 행으로 두고, 서버 시각·서버 스레드·커넥션 번호·기록 종류·p6spy가 기록한 실행 시간·SQL로 나눈다.
  2_요청별정리 : 기록을 요청(JMeter 요청 이름 C-2~C-7)별로 묶고, 요청 안에서는 트랜잭션별로 묶는다.
                 SQL에서 SQL 종류·대상 테이블·WHERE 조건·저장한 값을 꺼내 열로 둔다.
  3_시간값     : 두 자동 처리(C-3 모의 배송 완료, C-5 확인 기간 만료 자동 완료)의 SQL과 커밋을 한 표에 서버 시각순으로 둔다.
                 오른쪽에 서버 로그에 남은 두 자동 처리의 실행 기록(넘긴 now, 처리 건수)을 둔다.
각 시트 아래에는 열마다 무슨 데이터인지 적은 열 설명 표를 둔다.

p6spy 기록에는 서버 스레드 이름만 있고 어느 요청의 기록인지는 없다. 그래서 1번 스크립트로 뽑은 서버 로그(시각추출 파일)를 함께 읽어,
각 서버 스레드가 그 시각에 어느 요청을 처리하고 있었는지 알아낸다.
- 로그인한 회원의 요청은 「요청 처리 시작」 기록으로 시작을 찾는다(세부 작업 02의 실행02 스크립트와 같은 규칙).
- 테스트 지원 API(/test-support/jobs/...)의 요청은 보안 필터를 거치지 않아 「요청 처리 시작」 기록이 없다. 대신 요청이 끝날 때 남는
  「[테스트 자동 처리 실행]」 기록으로 그 요청을 찾고, 같은 스레드에서 그 앞에 남은 p6spy 기록을 그 요청의 기록으로 본다.

입력 파일은 읽기만 한다. 출력 파일이 이미 있으면 덮어쓰지 않고 중단한다.
실행 방법은 같은 폴더의 4_DB로그_엑셀정리_실행방법.txt를 참고한다.
"""
import argparse
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 서버 로그 줄의 형식
# 2026-10-20T11:20:03.750+09:00  INFO 1334666 --- [ChatService] [http-nio-8186-exec-10] p6spy    : #1792462803750 | took 5ms | statement | connection 11| url jdbc:oracle:thin:@127.0.0.1:1521:xe
LOG_LINE = re.compile(
    r'^(?P<time>\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\.\d{3}[+-]\d{2}:\d{2})\s+'
    r'(?P<level>[A-Z]+)\s+\d+\s+---\s+\[[^\]]*\]\s+\[\s*(?P<thread>[^\]]*?)\s*\]\s+'
    r'(?P<logger>\S+)\s*:\s(?P<message>.*)$'
)

# p6spy 기록 첫 줄의 메시지 부분. took는 p6spy가 기록한 실행 시간(ms), category는 기록 종류, connection은 커넥션 번호다.
P6SPY_MESSAGE = re.compile(r'^#\d+ \| took (?P<took>\d+)ms \| (?P<category>\w+) \| connection (?P<connection>\d+)\s*\|')

# 웹 애플리케이션 로그에서 요청을 나누고 이름을 붙일 때 쓰는 패턴
REQUEST_START = 'JWTAuthenticationFilter - 요청 처리 시작'
HTTP_METHOD = re.compile(r'HTTP Method: (?P<method>\S+), URL: \S+?/ChatService(?P<path>\S*)')
AUTH_USER_ID = re.compile(r'JWT 인증 성공: 사용자 ID = (?P<member>\S+)')
# 테스트 지원 API가 자동 처리를 마친 뒤 남기는 기록. 예) [테스트 자동 처리 실행] 모의 배송 완료 now=2026-10-20T02:20:03.701172Z, processed=1
JOB_END = re.compile(r'^\[테스트 자동 처리 실행\] (?P<job>.+?) now=(?P<now>\S+), processed=(?P<processed>\d+)$')

# 자동 처리 이름 → (테스트 지원 API 경로, JMeter 요청 이름)
JOBS = {
    '모의 배송 완료': ('/test-support/jobs/delivery-complete', 'C-3. 모의 배송 완료 실행'),
    '확인 기간 만료 자동 완료': ('/test-support/jobs/inspection-complete', 'C-5. 확인 기간 만료 자동 완료 실행(now: 상품 확인 기한)'),
}
LABEL_SALES = 'C-2. 판매자의 판매 목록 조회(주문 ID 얻기)'
LABEL_SELLER_DETAILS = ['C-4. 판매자의 주문 조회(배송 완료 뒤, 상품 확인 기한 얻기)', 'C-6. 판매자의 최종 주문 조회']
LABEL_BUYER_DETAIL = 'C-7. 구매자의 최종 주문 조회'
JOB_LABELS = [label for _, label in JOBS.values()]
REQUEST_ORDER = [LABEL_SALES, JOB_LABELS[0], LABEL_SELLER_DETAILS[0], JOB_LABELS[1], LABEL_SELLER_DETAILS[1], LABEL_BUYER_DETAIL,
                 '그 밖의 요청', '요청에 속하지 않은 기록']
ORDER_DETAIL_PATH = re.compile(r'^/api/orders/\d+$')

# 원본의 기록 종류를 엑셀에 적을 이름으로 바꾼다. 목록에 없는 종류는 원본 그대로 적는다.
CATEGORY_NAMES = {'statement': 'SQL 실행', 'commit': '커밋', 'rollback': '롤백'}

# 데이터를 바꾸는 SQL 종류. 트랜잭션 안에 하나라도 있으면 「데이터를 바꾼 트랜잭션」으로 본다.
CHANGE_KINDS = ('INSERT', 'UPDATE', 'DELETE')

# SQL 종류별로 대상 테이블 이름을 찾는 패턴
TABLE_PATTERNS = {
    'SELECT': re.compile(r'\bfrom\s+(\w+)', re.I),
    'INSERT': re.compile(r'^insert\s+into\s+(\w+)', re.I),
    'UPDATE': re.compile(r'^update\s+(\w+)', re.I),
    'DELETE': re.compile(r'^delete\s+from\s+(\w+)', re.I),
}

# 저장한 값 열에 적을 열. 이 실행에서 확인하는 값만 둔다.
# UPDATE PURCHASE_ORDER는 바뀌지 않은 열까지 모든 열을 다시 쓰므로, 전부 적으면 확인할 값을 찾기 어렵다.
KEY_COLUMNS = {
    'PURCHASE_ORDER': ['TRADE_STATUS', 'SHIPPING_STATUS', 'COMPLETION_CAUSE', 'DELIVERED_AT', 'INSPECTION_DEADLINE_AT', 'FINALIZED_AT'],
    'WALLET': ['MEMBER_ID', 'BALANCE'],
    'BALANCE_TRANSACTION': ['MEMBER_ID', 'TYPE', 'AMOUNT', 'BALANCE_AFTER', 'CREATED_AT'],
    'SHIPMENT': ['ORDER_ID', 'TRACKING_NUMBER', 'SHIPPED_AT'],
}

# Hibernate가 SQL의 열 이름 앞에 붙이는 테이블 별칭(po1_0. 같은 앞부분). WHERE 조건을 읽기 쉽게 하려고 지운다.
TABLE_ALIAS = re.compile(r'\b[a-z]+\d+_\d+\.')
# SELECT의 from 절에 적힌 「테이블 별칭」 쌍. 여러 테이블을 조회한 SQL에서 별칭을 테이블 이름으로 바꿀 때 쓴다.
FROM_TABLE = re.compile(r'(\w+)\s+([a-z]+\d+_\d+)\b')

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
TITLE_FONT = Font(bold=True, size=12)
WRAP = Alignment(wrap_text=True, vertical='top')


# ---------------------------------------------------------------------
# 입력 파일 읽기
# ---------------------------------------------------------------------
def read_p6spy_records(p6spy_path):
    """3번 스크립트로 뽑은 파일을 읽어 p6spy 기록 목록을 돌려준다.

    기록 하나는 시각 줄과 그 뒤의 두 줄(바인딩 전 SQL, 바인딩 후 SQL)로 이루어진다. 커밋 기록은 두 줄이 빈 줄과 ';'이다.
    """
    lines = p6spy_path.read_text(encoding='utf-8').splitlines()
    records = []
    for number, line in enumerate(lines, start=1):
        head = LOG_LINE.match(line)
        if head:
            message = P6SPY_MESSAGE.match(head.group('message'))
            if head.group('logger') != 'p6spy' or not message:
                sys.exit(f'p6spy 기록이 아닌 줄이 있습니다({number}번 줄). 3번 스크립트로 뽑은 파일인지 확인하세요.')
            records.append({
                'line': number,
                'head_text': line,
                'time': datetime.fromisoformat(head.group('time')),
                'thread': head.group('thread'),
                'connection': message.group('connection'),
                'category': message.group('category'),
                'took_ms': int(message.group('took')),
                'following': [],
            })
        elif records:
            records[-1]['following'].append(line)
    for record in records:
        following = record.pop('following')
        record['prepared'] = following[0] if following else ''
        record['bound'] = '\n'.join(following[1:])
        analyse_sql(record)
    return records


def new_request(thread, method='', path=''):
    return {'thread': thread, 'method': method, 'path': path, 'member': '', 'records': [], 'start_line': 0}


def assign_requests(log_path, records):
    """시각추출 파일을 읽어 p6spy 기록마다 시각추출 파일의 줄 번호와 요청을 정하고, 기록이 있는 요청 목록을 돌려준다.

    로그인한 회원의 요청: 같은 스레드의 「요청 처리 시작」 기록부터 그 스레드의 다음 「요청 처리 시작」 기록 전까지를 한 요청으로 본다.
    테스트 지원 API의 요청: 「요청 처리 시작」 기록이 없으므로, 처리 중인 요청이 없는 스레드에 남은 p6spy 기록을 모아 두었다가
    같은 스레드에 「[테스트 자동 처리 실행]」 기록이 나오면 모아 둔 기록을 그 요청의 기록으로 본다.
    """
    lines = log_path.read_text(encoding='utf-8').splitlines()
    requests = []
    open_request = {}   # 스레드 이름 → 그 스레드에서 지금 처리 중인 로그인 회원의 요청
    unclaimed = {}      # 스레드 이름 → 처리 중인 요청이 없을 때 남은 p6spy 기록
    pending = iter(records)
    next_record = next(pending, None)
    for number, line in enumerate(lines, start=1):
        match = LOG_LINE.match(line)
        if not match:
            continue
        thread = match.group('thread')
        if match.group('logger') == 'p6spy':
            # p6spy 추출 파일의 기록은 시각추출 파일과 같은 순서로 들어 있으므로, 첫 줄이 같은 기록을 앞에서부터 차례로 맞춘다.
            if next_record is not None and line == next_record['head_text']:
                next_record['log_line'] = number
                request = open_request.get(thread)
                if request is None:
                    unclaimed.setdefault(thread, []).append(next_record)
                else:
                    next_record['request'] = request
                next_record = next(pending, None)
            continue
        message = match.group('message')
        job = JOB_END.match(message)
        if job:
            if open_request.get(thread) is not None:
                sys.exit(f'{number}번 줄: 같은 스레드({thread})가 앞서 처리한 요청과 테스트 지원 API 요청을 구분할 수 없습니다.')
            path, label = JOBS.get(job.group('job'), ('', '그 밖의 요청'))
            request = new_request(thread, 'POST', path)
            request.update(label=label, job=job.group('job'), now=job.group('now'), processed=int(job.group('processed')))
            request['start_line'] = (unclaimed[thread][0]['log_line'] if unclaimed.get(thread) else number)
            for record in unclaimed.pop(thread, []):
                record['request'] = request
            requests.append(request)
            continue
        if message.startswith(REQUEST_START):
            open_request[thread] = new_request(thread)
            open_request[thread]['start_line'] = number
            requests.append(open_request[thread])
        request = open_request.get(thread)
        if request is None:
            continue
        method = HTTP_METHOD.search(message)
        if method and not request['method']:
            request['method'], request['path'] = method.group('method'), method.group('path')
        member = AUTH_USER_ID.search(message)
        if member and not request['member']:
            request['member'] = member.group('member')
    if next_record is not None:
        sys.exit('p6spy 추출 파일의 기록 가운데 시각추출 파일에서 찾지 못한 기록이 있습니다. 같은 실행에서 만든 두 파일인지 확인하세요.')

    label_requests([r for r in requests if 'label' not in r])
    # 처리 중인 요청을 찾지 못한 기록은 스레드별로 「요청에 속하지 않은 기록」으로 묶는다.
    outside = {}
    for thread, leftover in unclaimed.items():
        for record in leftover:
            record['request'] = outside.setdefault(thread, dict(new_request(thread), label='요청에 속하지 않은 기록'))
    for record in records:
        record['request']['records'].append(record)
    with_records = [r for r in requests if r['records']] + list(outside.values())
    return sorted(with_records, key=lambda r: (REQUEST_ORDER.index(r['label']), r['records'][0]['line']))


def label_requests(requests):
    """로그인한 회원의 요청마다 JMeter 요청 이름(C-2, C-4, C-6, C-7)을 붙인다.

    판매 목록 조회를 보낸 회원을 판매자로 본다. 판매자가 보낸 주문 조회는 먼저 시작한 것부터 C-4, C-6으로,
    판매자가 아닌 회원이 보낸 첫 주문 조회는 C-7로 정한다. 나머지는 「그 밖의 요청」이다.
    """
    seller = next((r['member'] for r in requests if r['path'].endswith('/orders/sales')), '')
    seller_details = iter(LABEL_SELLER_DETAILS)
    buyer_labelled = False
    for request in sorted(requests, key=lambda r: r['start_line']):
        path = request['path']
        if path.endswith('/orders/sales'):
            request['label'] = LABEL_SALES
        elif request['method'] == 'GET' and ORDER_DETAIL_PATH.match(path) and request['member'] == seller:
            request['label'] = next(seller_details, '그 밖의 요청')
        elif request['method'] == 'GET' and ORDER_DETAIL_PATH.match(path) and request['member'] and not buyer_labelled:
            request['label'] = LABEL_BUYER_DETAIL
            buyer_labelled = True
        else:
            request['label'] = '그 밖의 요청'


# ---------------------------------------------------------------------
# SQL과 트랜잭션 정리
# ---------------------------------------------------------------------
def split_values(text):
    """INSERT의 values 목록을 쉼표로 나눈다. 작은따옴표 안의 쉼표는 값의 일부이므로 나누지 않는다."""
    values, current, in_quote = [], '', False
    for char in text:
        if char == "'":
            in_quote = not in_quote
        if char == ',' and not in_quote:
            values.append(current.strip())
            current = ''
        else:
            current += char
    values.append(current.strip())
    return values


def find_changes(kind, table, sql):
    """UPDATE·INSERT가 쓴 값 가운데 KEY_COLUMNS에 있는 열의 값을 「열=값」 형식으로 돌려준다. 해당하지 않으면 빈 문자열을 돌려준다."""
    columns = KEY_COLUMNS.get(table.upper())
    if not columns:
        return ''
    found = []
    if kind == 'UPDATE':
        # set 뒤부터 마지막 where 앞까지가 바꾸는 값의 목록이다.
        set_part = re.split(r'\sset\s', sql, maxsplit=1, flags=re.I)[-1].rsplit(' where ', 1)[0]
        for column in columns:
            match = re.search(rf"(?<![\w.]){column}=('(?:[^']|'')*'|[^,\s]+)", set_part)
            if match:
                found.append(f'{column}={match.group(1)}')
    elif kind == 'INSERT':
        match = re.match(r'^insert\s+into\s+\w+\s*\((?P<columns>.*?)\)\s*values\s*\((?P<values>.*)\)\s*$', sql, re.I | re.S)
        if match:
            pairs = dict(zip([c.strip() for c in match.group('columns').split(',')], split_values(match.group('values'))))
            found = [f'{column}={pairs[column]}' for column in columns if column in pairs]
    return ', '.join(found)


def analyse_sql(record):
    """바인딩 후 SQL에서 SQL 종류, 대상 테이블, WHERE 조건, 저장한 값을 꺼낸다. 커밋·롤백 기록은 SQL 종류만 정한다."""
    if record['category'] != 'statement':
        record.update(sql_kind=CATEGORY_NAMES.get(record['category'], record['category']), table='', condition='', changes='')
        return
    sql = record['bound'].strip().rstrip(';').strip()
    kind = sql.split(None, 1)[0].upper() if sql else ''
    table_match = TABLE_PATTERNS[kind].search(sql) if kind in TABLE_PATTERNS else None
    table = table_match.group(1) if table_match else ''
    # 여러 테이블을 조회한 SELECT(예: 모의 배송 완료의 대상 주문 조회)는 대상 테이블을 모두 적고, WHERE 조건의 별칭을 테이블 이름으로 바꾼다.
    # 별칭을 지우기만 하면 「s1_0.ORDER_ID=po1_0.ORDER_ID」가 「ORDER_ID=ORDER_ID」가 되어 어느 테이블의 열인지 알 수 없다.
    aliases = {}
    if kind == 'SELECT':
        from_part = re.split(r'\sfrom\s', sql, maxsplit=1, flags=re.I)[-1]
        from_part = re.split(r'\s(?:where|order\s+by|fetch)\s', from_part, maxsplit=1, flags=re.I)[0]
        aliases = {alias: name for name, alias in FROM_TABLE.findall(from_part)}
        if len(aliases) > 1:
            table = ', '.join(aliases.values())
    # SELECT는 첫 where 뒤를, UPDATE·DELETE는 마지막 where 뒤를 조건으로 본다. SELECT 끝의 order by는 조건이 아니므로 뺀다.
    condition = ''
    if kind == 'SELECT' and re.search(r'\swhere\s', sql, re.I):
        condition = re.split(r'\swhere\s', sql, maxsplit=1, flags=re.I)[1]
        condition = re.split(r'\sorder\s+by\s', condition, maxsplit=1, flags=re.I)[0]
    elif kind in ('UPDATE', 'DELETE') and ' where ' in sql:
        condition = sql.rsplit(' where ', 1)[1]
    if len(aliases) > 1:
        condition = re.sub(r'\b([a-z]+\d+_\d+)\.', lambda m: aliases.get(m.group(1), m.group(1)) + '.', condition)
    else:
        condition = TABLE_ALIAS.sub('', condition)
    record.update(sql_kind=kind, table=table, condition=condition.strip(), changes=find_changes(kind, table, sql))


def split_transactions(request):
    """요청의 기록을 트랜잭션별로 묶는다. 같은 커넥션 번호의 기록을 커밋(또는 롤백) 기록까지 한 트랜잭션으로 본다."""
    transactions = []
    open_transaction = {}   # 커넥션 번호 → 아직 커밋·롤백 기록이 나오지 않은 트랜잭션
    for record in request['records']:
        transaction = open_transaction.get(record['connection'])
        if transaction is None:
            transaction = {'connection': record['connection'], 'records': [], 'end': None}
            open_transaction[record['connection']] = transaction
            transactions.append(transaction)
        transaction['records'].append(record)
        if record['category'] in ('commit', 'rollback'):
            transaction['end'] = record
            del open_transaction[record['connection']]
    for transaction in transactions:
        transaction['changed'] = any(r['sql_kind'] in CHANGE_KINDS for r in transaction['records'])
        nature = '데이터를 바꾼 트랜잭션' if transaction['changed'] else '조회만 한 트랜잭션'
        transaction['nature'] = nature if transaction['end'] else nature + '(커밋 기록 없음)'
    return transactions


def server_text(moment):
    """시각을 서버 로그와 같은 자릿수(밀리초까지)의 문자열로 바꾼다."""
    return moment.strftime('%Y-%m-%d %H:%M:%S.') + f'{moment.microsecond // 1000:03d}'


def summary(record):
    """3_시간값 시트의 칸에 적을 내용. 첫 줄은 SQL 종류와 대상 테이블, 그 아래는 저장한 값과 WHERE 조건이다."""
    if record['category'] != 'statement':
        return record['sql_kind']
    parts = [f"{record['sql_kind']} {record['table']}"]
    if record['changes']:
        parts.append(f"저장한 값: {record['changes']}")
    if record['condition']:
        parts.append(f"WHERE 조건: {record['condition']}")
    return '\n'.join(parts)


# ---------------------------------------------------------------------
# 엑셀 작성
# ---------------------------------------------------------------------
def write_table(sheet, start_row, start_col, headers, rows, title=None):
    """표 하나를 쓰고, 표 다음 빈 행 번호를 돌려준다."""
    row = start_row
    if title:
        sheet.cell(row=row, column=start_col, value=title).font = TITLE_FONT
        row += 1
    for offset, header in enumerate(headers):
        cell = sheet.cell(row=row, column=start_col + offset, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    for values in rows:
        row += 1
        for offset, value in enumerate(values):
            sheet.cell(row=row, column=start_col + offset, value=value).alignment = WRAP
    return row + 2


def set_widths(sheet, widths, start_col=1):
    """start_col 열부터 차례로 열 너비를 정한다."""
    for offset, width in enumerate(widths):
        sheet.column_dimensions[get_column_letter(start_col + offset)].width = width


def freeze_header(sheet, cell):
    """cell의 위쪽 행과 왼쪽 열을 고정하고, 처음 선택된 셀을 cell로 둔다.

    선택된 셀이 고정한 영역(A1)에 있으면 엑셀이 움직이는 영역을 1행부터 보여 주어, 머리글 행이 두 번 보인다.
    """
    sheet.freeze_panes = cell
    for selection in sheet.sheet_view.selection:
        if selection.pane == sheet.sheet_view.pane.activePane:
            selection.activeCell = cell
            selection.sqref = cell


TOOK_HEADER = 'p6spy가 기록한 실행 시간(ms)'
TOOK_NOTE = '원본의 took 값. 애플리케이션이 JDBC로 이 SQL이나 커밋을 보내고 결과를 받기까지 걸린 시간(밀리초)이다'


def build_sheet_original(workbook, records, p6spy_name, log_name):
    """1_원본 시트: p6spy 기록 하나를 한 행으로 둔다."""
    sheet = workbook.active
    sheet.title = '1_원본'
    headers = ['원본 줄 번호', '시각추출 파일 줄 번호', '서버 시각', '서버 스레드', '커넥션 번호', '기록 종류',
               TOOK_HEADER, 'SQL(바인딩 전)', 'SQL(바인딩 후)']
    rows = [[r['line'], r['log_line'], server_text(r['time']), r['thread'], int(r['connection']),
             CATEGORY_NAMES.get(r['category'], r['category']), r['took_ms'], r['prepared'], r['bound']] for r in records]
    next_row = write_table(sheet, 1, 1, headers, rows)
    freeze_header(sheet, 'A2')
    set_widths(sheet, [10, 12, 25, 22, 10, 10, 12, 70, 70])
    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['원본 줄 번호', f'{p6spy_name}에서 이 기록이 시작되는 줄의 번호'],
        ['시각추출 파일 줄 번호', f'{log_name}에서 이 기록이 시작되는 줄의 번호. 이 기록의 앞뒤에 남은 웹 애플리케이션 로그를 시각추출 파일에서 찾을 때 쓴다'],
        ['서버 시각', '서버가 로그 줄 맨 앞에 기록한 시각(한국 시각, 밀리초까지). p6spy는 SQL 실행이나 커밋이 끝난 뒤 기록하므로, 실행이 끝난 시각이다'],
        ['서버 스레드', '이 SQL이나 커밋을 보낸 서버 스레드의 이름'],
        ['커넥션 번호', 'p6spy가 DB 연결마다 붙인 번호. 서버는 트랜잭션을 시작할 때 DB 연결 하나를 빌려 커밋할 때까지 쓰므로, '
                   '번호가 같고 커밋 전까지 이어진 기록은 같은 트랜잭션에 속한다'],
        ['기록 종류', '원본의 statement는 「SQL 실행」, commit은 「커밋」, rollback은 「롤백」으로 적었다'],
        [TOOK_HEADER, TOOK_NOTE],
        ['SQL(바인딩 전)', '서버가 DB에 보낸 SQL. 값이 들어갈 자리는 ?로 적혀 있고, 실제 값은 SQL과 따로 보낸다. 커밋 기록은 비어 있다'],
        ['SQL(바인딩 후)', '바인딩 전 SQL의 ? 자리에 실제 값을 넣은 SQL(바인딩은 ? 자리에 실제 값을 넣는 것을 말한다). 어느 주문, 어느 회원, 얼마의 금액으로 실행했는지는 이 열에서 확인한다. 커밋 기록은 ;만 있다'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1, value=(
        '이 시트에는 p6spy 기록만 넣었다. 실행02에서는 웹 애플리케이션 로그를 엑셀로 정리하지 않았고, 요청을 나누는 데 필요한 기록만 시각추출 파일에서 읽었다. '
        '원본 첫 줄의 #숫자(서버 시각을 숫자로 다시 적은 값)와 url(접속한 DB 주소)은 새로운 정보가 없어 열로 두지 않았다.'))


def build_sheet_requests(workbook, requests):
    """2_요청별정리 시트: 기록을 요청별로 묶고, 요청 안에서는 트랜잭션별로 묶는다."""
    sheet = workbook.create_sheet('2_요청별정리')
    headers = ['요청 구분', '서버 스레드', '커넥션 번호', '트랜잭션 성격', '요청별 기록 순서', '원본 줄 번호', '서버 시각',
               'SQL 종류', '대상 테이블', 'WHERE 조건', '저장한 값', TOOK_HEADER, 'SQL(바인딩 후)']
    rows = []
    for request in requests:
        sequence = 0
        for transaction_index, transaction in enumerate(request['transactions']):
            for record_index, record in enumerate(transaction['records']):
                sequence += 1
                # 요청에 대한 값은 요청의 첫 행에만, 트랜잭션에 대한 값은 트랜잭션의 첫 행에만 적는다.
                # 모든 행에 적으면 같은 기록이 여러 번 들어간 것처럼 보인다.
                first_of_request = transaction_index == 0 and record_index == 0
                request_values = [request['label'], request['thread']] if first_of_request else ['', '']
                transaction_values = [int(transaction['connection']), transaction['nature']] if record_index == 0 else ['', '']
                rows.append(request_values + transaction_values + [
                    sequence, record['line'], server_text(record['time']), record['sql_kind'], record['table'],
                    record['condition'], record['changes'], record['took_ms'], record['bound'],
                ])
    next_row = write_table(sheet, 1, 1, headers, rows)
    freeze_header(sheet, 'B2')
    set_widths(sheet, [30, 22, 10, 22, 10, 10, 25, 10, 22, 40, 50, 12, 70])
    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['요청 구분', 'JMeter 요청 이름(결과.jtl의 label). C-2는 주문 ID를 얻는 판매 목록 조회, C-3은 모의 배송 완료, C-4는 배송 완료 뒤 판매자의 주문 조회, '
                 'C-5는 확인 기간 만료 자동 완료, C-6·C-7은 판매자·구매자의 최종 주문 조회다. C-1(멈춘 요청 목록 조회)은 SQL을 실행하지 않아 이 시트에 없다. '
                 '「요청에 속하지 않은 기록」은 처리 중인 요청을 찾지 못한 스레드의 기록이다'],
        ['서버 스레드', '그 요청을 처리한 서버 스레드의 이름'],
        ['커넥션 번호', '그 트랜잭션이 쓴 DB 연결의 번호'],
        ['트랜잭션 성격', 'UPDATE·INSERT·DELETE가 하나라도 있으면 「데이터를 바꾼 트랜잭션」, SELECT만 있으면 「조회만 한 트랜잭션」이다. '
                    '커밋·롤백 기록이 없으면 뒤에 「(커밋 기록 없음)」을 붙였다. C-3·C-5의 대상 주문 조회는 트랜잭션 밖에서 실행되어 커밋 기록이 없다'],
        ['요청별 기록 순서', '같은 요청이 남긴 p6spy 기록 가운데 몇 번째 기록인지를 나타낸다'],
        ['원본 줄 번호', '1_원본 시트의 원본 줄 번호와 같은 값'],
        ['서버 시각', '서버가 기록한 시각(한국 시각, 밀리초까지). SQL 실행이나 커밋이 끝난 시각이다'],
        ['SQL 종류', 'SQL의 첫 단어(SELECT, INSERT, UPDATE, DELETE). 커밋·롤백 기록은 「커밋」·「롤백」이다'],
        ['대상 테이블', 'SQL이 읽거나 바꾼 테이블. 여러 테이블을 조회한 SELECT는 from 절의 테이블을 모두 적었다'],
        ['WHERE 조건', 'SQL의 where 뒤에 적힌 조건. 열 이름 앞의 테이블 별칭(po1_0. 같은 앞부분)은 지웠고, '
                    '여러 테이블을 조회한 SELECT는 별칭을 테이블 이름으로 바꾸었다'],
        ['저장한 값', 'UPDATE·INSERT가 테이블에 저장한 값 가운데 이 실행에서 확인하는 열의 값. '
                 'PURCHASE_ORDER는 TRADE_STATUS·SHIPPING_STATUS·COMPLETION_CAUSE·DELIVERED_AT·INSPECTION_DEADLINE_AT·FINALIZED_AT, '
                 'WALLET은 MEMBER_ID·BALANCE, BALANCE_TRANSACTION은 MEMBER_ID·TYPE·AMOUNT·BALANCE_AFTER·CREATED_AT, SHIPMENT는 ORDER_ID·TRACKING_NUMBER·SHIPPED_AT이다. '
                 'UPDATE는 SET에 있는 열만 적으므로, WHERE 조건에 쓴 열(예: UPDATE WALLET의 MEMBER_ID)은 WHERE 조건 열에서 본다. 그 밖의 테이블은 비워 둔다'],
        [TOOK_HEADER, TOOK_NOTE],
        ['SQL(바인딩 후)', '1_원본 시트의 같은 열과 같은 값'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1, value=(
        '요청의 범위: 로그인한 회원의 요청은 같은 서버 스레드의 「요청 처리 시작」 기록부터 그 스레드의 다음 「요청 처리 시작」 기록 전까지를 한 요청으로 보았다. '
        '테스트 지원 API의 요청(C-3, C-5)은 「요청 처리 시작」 기록이 없어, 같은 스레드에 「[테스트 자동 처리 실행]」 기록이 남기 전까지 남은 기록을 그 요청의 기록으로 보았다. '
        '트랜잭션의 범위: 같은 요청 안에서 커넥션 번호가 같은 기록을 커밋(또는 롤백) 기록까지 한 트랜잭션으로 보았다. '
        '요청 구분과 서버 스레드는 각 요청의 첫 행에만, 커넥션 번호와 트랜잭션 성격은 각 트랜잭션의 첫 행에만 적었다. '
        '그 아래 빈칸은 바로 위에 적힌 값과 같다.'))


def build_sheet_times(workbook, requests):
    """3_시간값 시트: 두 자동 처리(C-3, C-5)의 SQL과 커밋을 한 표에 서버 시각순으로 두고, 오른쪽에 자동 처리의 실행 기록을 둔다."""
    sheet = workbook.create_sheet('3_시간값')
    jobs = [r for r in requests if r['label'] in JOB_LABELS]
    rows = []
    for request in jobs:
        for transaction in request['transactions']:
            for record in transaction['records']:
                rows.append((record['time'], record['line'], [
                    server_text(record['time']), request['label'].split('.')[0], int(transaction['connection']),
                    transaction['nature'], summary(record)]))
    rows.sort(key=lambda item: (item[0], item[1]))
    headers = ['서버 시각', '요청 구분', '커넥션 번호', '트랜잭션 성격', 'SQL이나 커밋']
    next_row = write_table(sheet, 1, 1, headers, [values for _, _, values in rows])
    freeze_header(sheet, 'A2')
    set_widths(sheet, [25, 10, 10, 26, 80])

    job_rows = [[r['label'], r['path'], r['now'], r['processed'], server_text(r['records'][-1]['time'])] for r in jobs]
    write_table(sheet, 1, 7, ['요청 구분', '테스트 지원 API', '넘긴 now(UTC)', '처리 건수', '마지막 p6spy 기록 시각'], job_rows,
                title='자동 처리 실행 기록(시각추출 파일의 「[테스트 자동 처리 실행]」 기록)')
    sheet.cell(row=len(job_rows) + 4, column=7, value=(
        '넘긴 now와 처리 건수는 테스트 지원 API가 자동 처리를 마친 뒤 서버 로그에 남긴 값이다. '
        'now를 넘기지 않은 요청(C-3)은 서버의 현재 시각이 적혀 있다.'))
    set_widths(sheet, [40, 36, 30, 10, 25], start_col=7)

    note_row = write_table(sheet, next_row, 1, ['열', '내용'], [
        ['서버 시각', '서버가 기록한 시각(한국 시각, 밀리초까지). SQL 실행이나 커밋이 끝난 시각이다'],
        ['요청 구분', 'C-3(모의 배송 완료) 또는 C-5(확인 기간 만료 자동 완료)'],
        ['커넥션 번호', '그 SQL이나 커밋을 보낸 DB 연결의 번호'],
        ['트랜잭션 성격', '2_요청별정리 시트의 같은 열과 같은 값. 커밋 기록이 없는 것은 대상 주문 조회, 데이터를 바꾼 트랜잭션은 주문 하나의 처리다'],
        ['SQL이나 커밋', '첫 줄은 SQL 종류와 대상 테이블, 그 아래는 저장한 값과 WHERE 조건이다. 커밋 기록은 「커밋」이다'],
    ], title='열 설명')
    sheet.cell(row=note_row, column=1, value=(
        '이 표에는 두 자동 처리의 기록만 넣었다. 주문 조회 요청(C-2, C-4, C-6, C-7)의 기록은 2_요청별정리 시트에 있다. '
        '모든 시각이 같은 서버 시계로 기록되었으므로 행의 앞뒤를 그대로 비교할 수 있다. 서버 시각이 같은 행은 원본 줄 순서대로 두었다.'))


def main():
    parser = argparse.ArgumentParser(description='p6spy 기록(DB 쪽 로그)을 엑셀 파일로 정리한다.')
    parser.add_argument('--p6spy', required=True, help='3번 스크립트로 뽑은 p6spy 기록 파일')
    parser.add_argument('--log', required=True, help='1번 스크립트로 뽑은 서버 로그 파일(시각추출 파일). 요청을 알아내는 데만 쓴다')
    parser.add_argument('--output', required=True, help='만들 엑셀 파일. 같은 이름의 파일이 있으면 중단한다')
    args = parser.parse_args()

    p6spy_path, log_path, output = Path(args.p6spy), Path(args.log), Path(args.output)
    if output.exists():
        sys.exit(f'출력 파일이 이미 있습니다. 다른 이름을 지정하세요: {output}')

    records = read_p6spy_records(p6spy_path)
    requests = assign_requests(log_path, records)
    for request in requests:
        request['transactions'] = split_transactions(request)

    workbook = Workbook()
    build_sheet_original(workbook, records, p6spy_path.name, log_path.name)
    build_sheet_requests(workbook, requests)
    build_sheet_times(workbook, requests)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    kinds = Counter(CATEGORY_NAMES.get(r['category'], r['category']) for r in records)
    print(f'p6spy 기록: {len(records)}건 (' + ', '.join(f'{kind} {count}건' for kind, count in kinds.items()) + ')')
    for request in requests:
        extra = f", now={request['now']}, 처리 건수 {request['processed']}" if 'now' in request else ''
        print(f"  {request['label']}: {request['thread']}, 기록 {len(request['records'])}건{extra}")
        for transaction in request['transactions']:
            end = server_text(transaction['end']['time'])[11:] if transaction['end'] else '없음'
            print(f"    커넥션 {transaction['connection']}: {transaction['nature']}, 기록 {len(transaction['records'])}건, 커밋 {end}")
    print(f'저장한 파일: {output}')


if __name__ == '__main__':
    main()
